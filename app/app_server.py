import os
import logging
from datetime import datetime
from flask import Flask, request, jsonify, send_file
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = Flask(__name__)

progress_data = {'status': 'idle', 'progress': 0, 'stage': '', 'message': '', 'elapsed_time': 0, 'estimated_time': 0}

@app.route('/')
def index():
    return '''<!DOCTYPE html><html><head><title>Bill vs Pay</title><style>
    body{font-family:Arial;max-width:900px;margin:50px auto;padding:20px;background:#f5f5f5;}
    .container{background:white;padding:30px;border-radius:8px;box-shadow:0 2px 8px rgba(0,0,0,0.1);}
    h1{color:#1F5F99;margin:0;}label{display:block;margin:15px 0 5px 0;font-weight:bold;}
    input[type="file"]{display:block;margin:5px 0 15px 0;padding:8px;width:100%;border:1px solid #ddd;border-radius:4px;}
    button{background:#1F5F99;color:white;padding:12px;border:none;border-radius:4px;cursor:pointer;width:100%;font-size:16px;}
    button:hover{background:#146c43;}button:disabled{background:#999;cursor:not-allowed;}
    .status{margin:20px 0;padding:15px;background:#e3f2fd;border-radius:4px;display:none;}
    .status.active{display:block;}.error{background:#ffebee;color:#c62828;}.success{background:#e8f5e9;color:#2e7d32;}
    .progress-section{display:none;margin:20px 0;padding:20px;background:#f0f4f8;border-radius:5px;}
    .progress-section.active{display:block;}.progress-bar{background:#ddd;height:30px;border-radius:4px;overflow:hidden;margin:10px 0;}
    .progress-fill{background:linear-gradient(90deg,#1F5F99,#4472C4);height:100%;width:0%;display:flex;align-items:center;justify-content:center;color:white;font-weight:bold;font-size:12px;}
    </style></head><body><div class="container"><h1>Bill vs Pay Reconciliation</h1><p>Upload payroll and billing exports</p>
    <div><label>Payroll (Excel):</label><input type="file" id="payroll" accept=".xlsx" required>
    <label>Billing (Excel):</label><input type="file" id="billing" accept=".xlsx" required>
    <button onclick="reconcile()">Reconcile Now</button></div>
    <div id="progress" class="progress-section"><div class="progress-bar"><div id="fill" class="progress-fill" style="width:0%">0%</div></div>
    <div id="msg"></div></div>
    <div id="status" class="status"></div></div>
    <script>
    async function reconcile(){
        const payroll=document.getElementById('payroll').files[0];
        const billing=document.getElementById('billing').files[0];
        if(!payroll||!billing){alert('Select both files');return;}
        
        document.getElementById('progress').classList.add('active');
        document.getElementById('status').classList.remove('active');
        
        const fd=new FormData();fd.append('payroll',payroll);fd.append('billing',billing);
        
        let poll=setInterval(async()=>{
            const r=await fetch('/progress');
            const d=await r.json();
            document.getElementById('fill').style.width=d.progress+'%';
            document.getElementById('fill').textContent=d.progress+'%';
            document.getElementById('msg').innerHTML='<strong>'+d.stage+'</strong><br>'+d.message;
        },500);
        
        try{
            const r=await fetch('/reconcile',{method:'POST',body:fd});
            if(r.ok){
                clearInterval(poll);
                const blob=await r.blob();
                const url=URL.createObjectURL(blob);
                const a=document.createElement('a');
                a.href=url;
                a.download='reconciliation.xlsx';
                a.click();
                document.getElementById('status').className='status active success';
                document.getElementById('status').textContent='✓ Success! Report downloaded.';
                document.getElementById('progress').classList.remove('active');
            }else{
                clearInterval(poll);
                const t=await r.text();
                document.getElementById('status').className='status active error';
                document.getElementById('status').textContent='✗ Error: '+t;
            }
        }catch(e){
            clearInterval(poll);
            document.getElementById('status').className='status active error';
            document.getElementById('status').textContent='✗ Error: '+e.message;
        }
    }
    </script></body></html>'''

@app.route('/progress')
def progress():
    return jsonify(progress_data)

@app.route('/reconcile', methods=['POST'])
def reconcile():
    global progress_data
    try:
        progress_data = {'status':'processing','progress':5,'stage':'Loading','message':'Reading files...','elapsed_time':0,'estimated_time':0}
        
        pf = request.files.get('payroll')
        bf = request.files.get('billing')
        if not pf or not bf:
            return 'Missing files', 400
        
        progress_data = {'status':'processing','progress':20,'stage':'Loading','message':'Loading payroll...','elapsed_time':0,'estimated_time':0}
        payroll_df = pd.read_excel(pf)
        
        progress_data = {'status':'processing','progress':40,'stage':'Loading','message':'Loading billing...','elapsed_time':0,'estimated_time':0}
        billing_df = pd.read_excel(bf)
        
        progress_data = {'status':'processing','progress':60,'stage':'Reconciling','message':'Processing...','elapsed_time':0,'estimated_time':0}
        
        from reconciliation import BillPayReconciler
        reconciler = BillPayReconciler(payroll_df, billing_df)
        results = reconciler.reconcile()
        
        progress_data = {'status':'processing','progress':85,'stage':'Excel','message':'Creating report...','elapsed_time':0,'estimated_time':0}
        
        wb = Workbook()
        ws = wb.active
        ws.title = 'Summary'
        ws['A1'] = 'Bill vs Pay Reconciliation Report'
        ws['A1'].font = Font(size=14, bold=True, color="FFFFFF")
        ws['A1'].fill = PatternFill(start_color="1F5F99", end_color="1F5F99", fill_type="solid")
        
        summary = results['summary']
        ws['A3'] = 'Total Matches'
        ws['B3'] = summary['total_matches']
        ws['A4'] = 'Match Rate %'
        ws['B4'] = f"{summary['match_rate_payroll']:.2f}%"
        ws['A5'] = 'Matched Amount'
        ws['B5'] = f"${summary['matched_payroll_amount']:,.2f}"
        ws['A6'] = 'Unmatched Payroll'
        ws['B6'] = f"${summary['unmatched_payroll_amount']:,.2f}"
        ws['A7'] = 'Unmatched Billing'
        ws['B7'] = f"${summary['unmatched_billing_amount']:,.2f}"
        
        if len(results['matches']) > 0:
            ws_m = wb.create_sheet('Matches')
            matches_df = results['matches']
            for col_idx, col_name in enumerate(matches_df.columns, 1):
                ws_m.cell(row=1, column=col_idx, value=col_name)
            for row_idx, row_data in enumerate(matches_df.itertuples(), 2):
                for col_idx, value in enumerate(row_data[1:], 1):
                    ws_m.cell(row=row_idx, column=col_idx, value=value)
        
        output_path = f'/tmp/reconciliation_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
        wb.save(output_path)
        
        progress_data = {'status':'complete','progress':100,'stage':'Done','message':'Report ready','elapsed_time':0,'estimated_time':0}
        
        return send_file(output_path, as_attachment=True, download_name=f"reconciliation_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx", mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    
    except Exception as e:
        logger.error(f"Error: {e}", exc_info=True)
        return f'Error: {str(e)}', 500

@app.route('/health')
def health():
    return jsonify({'status': 'ok'})

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=int(os.environ.get('PORT', 8765)), debug=False)
