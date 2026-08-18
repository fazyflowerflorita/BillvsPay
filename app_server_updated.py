"""
Updated Flask Server for Bill vs Pay Reconciliation
Integrates the new reconciliation_engine_v2.py
"""

from flask import Flask, request, jsonify, send_file, render_template
import pandas as pd
import logging
from pathlib import Path
from datetime import datetime
import io
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Import the new reconciliation engine
from reconciliation_engine_v2 import BillPayReconciler

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = Flask(__name__)

# Configuration
UPLOAD_FOLDER = Path('uploads')
OUTPUT_FOLDER = Path('output')
MAX_FILE_SIZE = 50 * 1024 * 1024  # 50MB

UPLOAD_FOLDER.mkdir(exist_ok=True)
OUTPUT_FOLDER.mkdir(exist_ok=True)


class ReconciliationApp:
    """Main reconciliation application"""
    
    def __init__(self):
        self.reconciler = None
        self.results = None
    
    def validate_files(self, payroll_file, billing_file):
        """Validate uploaded files"""
        if not payroll_file or not billing_file:
            raise ValueError("Both payroll and billing files are required")
        
        if not payroll_file.filename.endswith('.xlsx'):
            raise ValueError("Payroll file must be .xlsx format")
        
        if not billing_file.filename.endswith('.xlsx'):
            raise ValueError("Billing file must be .xlsx format")
        
        if payroll_file.content_length > MAX_FILE_SIZE:
            raise ValueError(f"Payroll file exceeds {MAX_FILE_SIZE / 1024 / 1024}MB limit")
        
        if billing_file.content_length > MAX_FILE_SIZE:
            raise ValueError(f"Billing file exceeds {MAX_FILE_SIZE / 1024 / 1024}MB limit")
    
    def load_files(self, payroll_file, billing_file):
        """Load Excel files with error handling"""
        try:
            # Save temporarily
            payroll_path = UPLOAD_FOLDER / f"payroll_{datetime.now().timestamp()}.xlsx"
            billing_path = UPLOAD_FOLDER / f"billing_{datetime.now().timestamp()}.xlsx"
            
            payroll_file.save(payroll_path)
            billing_file.save(billing_path)
            
            logger.info(f"Loading payroll from {payroll_path}")
            logger.info(f"Loading billing from {billing_path}")
            
            # Load with error handling
            try:
                payroll_df = pd.read_excel(payroll_path)
            except Exception as e:
                raise ValueError(f"Error reading payroll file: {e}")
            
            try:
                billing_df = pd.read_excel(billing_path)
            except Exception as e:
                raise ValueError(f"Error reading billing file: {e}")
            
            logger.info(f"Loaded {len(payroll_df)} payroll rows and {len(billing_df)} billing rows")
            
            return payroll_df, billing_df
        
        except Exception as e:
            logger.error(f"File loading error: {e}")
            raise
    
    def run_reconciliation(self, payroll_df, billing_df):
        """Run the reconciliation process"""
        try:
            logger.info("Starting reconciliation...")
            
            self.reconciler = BillPayReconciler(payroll_df, billing_df)
            self.results = self.reconciler.reconcile()
            
            logger.info("Reconciliation completed successfully")
            return self.results
        
        except Exception as e:
            logger.error(f"Reconciliation error: {e}")
            raise
    
    def generate_excel_output(self):
        """Generate comprehensive Excel output workbook"""
        try:
            if not self.results:
                raise ValueError("No reconciliation results available")
            
            wb = Workbook()
            wb.remove(wb.active)  # Remove default sheet
            
            # 1. Summary Tab
            self._create_summary_tab(wb)
            
            # 2. Match Details Tab
            self._create_match_details_tab(wb)
            
            # 3. Unmatched Payroll Tab
            self._create_unmatched_payroll_tab(wb)
            
            # 4. Unmatched Billing Tab
            self._create_unmatched_billing_tab(wb)
            
            # 5. Exception Summary Tab
            self._create_exception_summary_tab(wb)
            
            # Save to file
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_path = OUTPUT_FOLDER / f"reconciliation_{timestamp}.xlsx"
            wb.save(output_path)
            
            logger.info(f"Excel report generated: {output_path}")
            return output_path
        
        except Exception as e:
            logger.error(f"Error generating Excel output: {e}")
            raise
    
    def _create_summary_tab(self, wb):
        """Create summary tab"""
        ws = wb.create_sheet('Summary', 0)
        
        summary = self.results['summary']
        
        # Title
        ws['A1'] = 'Bill vs Pay Reconciliation Summary'
        ws['A1'].font = Font(size=14, bold=True, color="FFFFFF")
        ws['A1'].fill = PatternFill(start_color="1F5F99", end_color="1F5F99", fill_type="solid")
        ws.merge_cells('A1:B1')
        
        # Data
        data = [
            ('', ''),
            ('Report Date', datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
            ('', ''),
            ('OVERALL RESULTS', ''),
            ('Total Payroll Rows', summary['total_payroll_rows']),
            ('Total Billing Rows', summary['total_billing_rows']),
            ('Total Matched', summary['total_matches']),
            ('Match Rate (Payroll %)', f"{summary['match_rate_payroll']:.2f}%"),
            ('', ''),
            ('AMOUNT RECONCILIATION', ''),
            ('Total Payroll Amount', f"${summary['payroll_amount_total']:,.2f}"),
            ('Total Billing Amount', f"${summary['billing_amount_total']:,.2f}"),
            ('Matched Payroll Amount', f"${summary['matched_payroll_amount']:,.2f}"),
            ('Matched Billing Amount', f"${summary['matched_billing_amount']:,.2f}"),
            ('Unmatched Payroll Amount', f"${summary['unmatched_payroll_amount']:,.2f}"),
            ('Unmatched Billing Amount', f"${summary['unmatched_billing_amount']:,.2f}"),
            ('', ''),
            ('UNMATCHED SUMMARY', ''),
            ('Paid But Not Billed (Rows)', summary['unmatched_payroll_rows']),
            ('Billed But Not Paid (Rows)', summary['unmatched_billing_rows']),
        ]
        
        for idx, (label, value) in enumerate(data, 2):
            ws[f'A{idx}'] = label
            ws[f'B{idx}'] = value
            
            if label in ['OVERALL RESULTS', 'AMOUNT RECONCILIATION', 'UNMATCHED SUMMARY']:
                ws[f'A{idx}'].font = Font(bold=True, color="FFFFFF")
                ws[f'A{idx}'].fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                ws.merge_cells(f'A{idx}:B{idx}')
        
        ws.column_dimensions['A'].width = 30
        ws.column_dimensions['B'].width = 20
    
    def _create_match_details_tab(self, wb):
        """Create detailed match results tab"""
        ws = wb.create_sheet('Match Details')
        
        matches_df = self.results['matches']
        
        if len(matches_df) == 0:
            ws['A1'] = 'No matches found'
            return
        
        # Headers
        headers = [
            'Employee ID', 'Employee Name', 'Category', 'Source',
            'Payroll Amount', 'Billing Amount', 'Difference',
            'Payroll Qty', 'Billing Qty', 'Qty Diff',
            'Status/Comment', 'Payroll Date', 'Billing Date'
        ]
        
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="1F5F99", end_color="1F5F99", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        # Data rows with conditional formatting
        for row_idx, row in enumerate(matches_df.itertuples(), 2):
            ws.cell(row=row_idx, column=1).value = row.employee_id
            ws.cell(row=row_idx, column=2).value = row.employee_name
            ws.cell(row=row_idx, column=3).value = row.category
            ws.cell(row=row_idx, column=4).value = row.source
            ws.cell(row=row_idx, column=5).value = row.payroll_amount
            ws.cell(row=row_idx, column=6).value = row.billing_amount
            ws.cell(row=row_idx, column=7).value = row.amount_difference
            ws.cell(row=row_idx, column=8).value = row.payroll_qty
            ws.cell(row=row_idx, column=9).value = row.billing_qty
            ws.cell(row=row_idx, column=10).value = row.qty_difference
            ws.cell(row=row_idx, column=11).value = row.comment
            ws.cell(row=row_idx, column=12).value = row.payroll_date
            ws.cell(row=row_idx, column=13).value = row.billing_date
            
            # Color coding based on comment
            comment = row.comment
            if 'Tallied' in comment:
                fill_color = "00B050"  # Green
            elif 'Difference' in comment:
                fill_color = "FFC000"  # Orange
            else:
                fill_color = "FFFFFF"  # White
            
            for col in range(1, 14):
                ws.cell(row=row_idx, column=col).fill = PatternFill(
                    start_color=fill_color, end_color=fill_color, fill_type="solid"
                )
        
        # Set column widths
        widths = [12, 20, 15, 12, 15, 15, 12, 12, 12, 10, 30, 12, 12]
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(idx)].width = width
    
    def _create_unmatched_payroll_tab(self, wb):
        """Create unmatched payroll tab"""
        ws = wb.create_sheet('Unmatched Payroll')
        
        unmatched = self.results['unmatched_payroll']
        
        if len(unmatched) == 0:
            ws['A1'] = 'All payroll records matched!'
            return
        
        # Headers
        headers = ['Employee ID', 'Employee Name', 'Category', 'Amount', 'Quantity', 'Date', 'Status']
        
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="FF0000", end_color="FF0000", fill_type="solid")
        
        # Data
        for row_idx, row in enumerate(unmatched.itertuples(), 2):
            ws.cell(row=row_idx, column=1).value = row._0 if hasattr(row, '_0') else 'N/A'
            ws.cell(row=row_idx, column=2).value = row._1 if hasattr(row, '_1') else 'N/A'
            # ... add other columns
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20
    
    def _create_unmatched_billing_tab(self, wb):
        """Create unmatched billing tab"""
        ws = wb.create_sheet('Unmatched Billing')
        
        unmatched = self.results['unmatched_billing']
        
        if len(unmatched) == 0:
            ws['A1'] = 'All billing records matched!'
            return
        
        # Similar structure to unmatched payroll
        headers = ['Employee ID', 'Employee Name', 'Category', 'Amount', 'Quantity', 'Date', 'Status']
        
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid")
        
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20
    
    def _create_exception_summary_tab(self, wb):
        """Create exception summary tab"""
        ws = wb.create_sheet('Exception Summary')
        
        matches_df = self.results['matches']
        
        # Summary by comment type
        ws['A1'] = 'Exception Summary'
        ws['A1'].font = Font(size=12, bold=True)
        
        if len(matches_df) > 0:
            comment_counts = matches_df['comment'].value_counts()
            
            ws['A3'] = 'Status'
            ws['B3'] = 'Count'
            
            for idx, (comment, count) in enumerate(comment_counts.items(), 4):
                ws[f'A{idx}'] = comment
                ws[f'B{idx}'] = count
        
        ws.column_dimensions['A'].width = 50
        ws.column_dimensions['B'].width = 15


# Flask Routes
app_instance = ReconciliationApp()


@app.route('/', methods=['GET'])
def index():
    """Home page with upload form"""
    return render_template('upload.html')


@app.route('/reconcile', methods=['POST'])
def reconcile():
    """Main reconciliation endpoint"""
    try:
        # Get files
        payroll_file = request.files.get('payroll')
        billing_file = request.files.get('billing')
        
        # Validate
        app_instance.validate_files(payroll_file, billing_file)
        
        # Load
        payroll_df, billing_df = app_instance.load_files(payroll_file, billing_file)
        
        # Reconcile
        results = app_instance.run_reconciliation(payroll_df, billing_df)
        
        # Generate output
        output_path = app_instance.generate_excel_output()
        
        # Return file
        return send_file(
            output_path,
            as_attachment=True,
            download_name=f"reconciliation_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        )
    
    except ValueError as e:
        logger.warning(f"Validation error: {e}")
        return jsonify({'error': str(e)}), 400
    
    except Exception as e:
        logger.error(f"Reconciliation error: {e}")
        return jsonify({'error': f"Reconciliation failed: {str(e)}"}), 500


@app.route('/summary', methods=['GET'])
def get_summary():
    """Return current reconciliation summary as JSON"""
    if not app_instance.results:
        return jsonify({'error': 'No reconciliation results available'}), 404
    
    summary = app_instance.results['summary']
    return jsonify(summary)


@app.errorhandler(413)
def request_entity_too_large(error):
    return jsonify({'error': 'File too large (max 50MB)'}), 413


@app.errorhandler(500)
def internal_error(error):
    logger.error(f"Internal server error: {error}")
    return jsonify({'error': 'Internal server error'}), 500


if __name__ == '__main__':
    app.run(
        host=os.environ.get('HOST', '127.0.0.1'),
        port=int(os.environ.get('PORT', 8765)),
        debug=False
    )
