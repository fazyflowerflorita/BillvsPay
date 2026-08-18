"""
Bill vs Pay Reconciliation Web Server - Minimal & Robust
Flask application with error recovery and progress tracking
"""

import os
import logging
import time
from datetime import datetime
from pathlib import Path

from flask import Flask, request, jsonify, send_file

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

# Create Flask app
app = Flask(__name__)

# Configuration
UPLOAD_FOLDER = Path('uploads')
OUTPUT_FOLDER = Path('output')
MAX_FILE_SIZE = 50 * 1024 * 1024

UPLOAD_FOLDER.mkdir(exist_ok=True)
OUTPUT_FOLDER.mkdir(exist_ok=True)

# Global progress tracking
current_progress = {
    'status': 'idle',
    'stage': '',
    'progress': 0,
    'message': '',
    'estimated_time': 0,
    'elapsed_time': 0
}

start_time = None


def update_progress(stage, progress, message=''):
    """Update progress globally"""
    global current_progress, start_time
    
    if start_time is None:
        start_time = time.time()
    
    elapsed = time.time() - start_time
    
    if progress > 5:
        estimated_total = (elapsed / progress) * 100
        estimated_remaining = estimated_total - elapsed
    else:
        estimated_remaining = 0
    
    current_progress = {
        'status': 'processing',
        'stage': stage,
        'progress': min(int(progress), 100),
        'message': message,
        'estimated_time': max(0, int(estimated_remaining)),
        'elapsed_time': int(elapsed)
    }
    
    logger.info(f"{stage} - {progress}% - {message}")


@app.route('/', methods=['GET'])
def index():
    """Home page with file upload"""
    return '''
    <!DOCTYPE html>
    <html>
    <head>
        <title>Bill vs Pay Reconciliation</title>
        <meta charset="utf-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">
        <style>
            * { box-sizing: border-box; }
            body { 
                font-family: Arial, sans-serif; 
                max-width: 900px; 
                margin: 0 auto;
                background: #f5f5f5;
                padding: 20px;
            }
            .container { 
                background: white; 
                padding: 30px; 
                border-radius: 8px;
                box-shadow: 0 2px 8px rgba(0,0,0,0.1);
            }
            h1 { 
                color: #1F5F99;
                margin: 0 0 10px 0;
            }
            p {
                color: #666;
                margin: 0 0 20px 0;
            }
            .upload-section { 
                background: #f9f9f9; 
                padding: 20px; 
                border-radius: 5px;
                margin-bottom: 20px;
            }
            label { 
                display: block; 
                margin: 15px 0 5px 0;
                font-weight: bold;
                color: #333;
            }
            input[type="file"] { 
                display: block; 
                margin: 5px 0 15px 0;
                padding: 8px;
                border: 1px solid #ddd;
                border-radius: 4px;
                width: 100%;
            }
            button { 
                background: #1F5F99; 
                color: white; 
                padding: 12px 30px; 
                border: none; 
                border-radius: 4px; 
                cursor: pointer;
                font-size: 16px;
                width: 100%;
            }
            button:hover { background: #146c43; }
            button:disabled { background: #999; cursor: not-allowed; }
            
            .progress-section {
                display: none;
                margin-top: 20px;
                padding: 20px;
                background: #f0f4f8;
                border-radius: 5px;
                border-left: 4px solid #1F5F99;
            }
            .progress-section.active { display: block; }
            
            .progress-bar {
                background: #ddd;
                height: 30px;
                border-radius: 4px;
                overflow: hidden;
                margin: 10px 0;
            }
            .progress-fill {
                background: linear-gradient(90deg, #1F5F99, #4472C4);
                height: 100%;
                width: 0%;
                transition: width 0.3s ease;
                display: flex;
                align-items: center;
                justify-content: center;
                color: white;
                font-weight: bold;
                font-size: 12px;
            }
            
            .progress-info {
                display: grid;
                grid-template-columns: 1fr 1fr;
                gap: 10px;
                margin-top: 10px;
            }
            .info-item {
                background: white;
                padding: 10px;
                border-radius: 4px;
            }
            .info-label {
                font-weight: bold;
                color: #666;
                font-size: 12px;
            }
            .info-value {
                color: #1F5F99;
                font-size: 18px;
                margin-top: 5px;
            }
            
            .status-message {
                margin-top: 10px;
                padding: 10px;
                background: white;
                border-radius: 4px;
                color: #333;
                min-height: 20px;
            }
            
            .error-section {
                display: none;
                margin-top: 20px;
                padding: 15px;
                background: #ffebee;
                border-left: 4px solid #ff0000;
                border-radius: 4px;
                color: #c62828;
            }
            .error-section.active { display: block; }
            
            .success-section {
                display: none;
                margin-top: 20px;
                padding: 15px;
                background: #e8f5e9;
                border-left: 4px solid #4caf50;
                border-radius: 4px;
                color: #2e7d32;
            }
            .success-section.active { display: block; }
        </style>
    </head>
    <body>
        <div class="container">
            <h1>Bill vs Pay Reconciliation</h1>
            <p>Upload payroll and billing exports for automated reconciliation</p>
            
            <div class="upload-section">
                <form id="reconcileForm">
                    <label>Payroll Report (Excel)</label>
                    <input type="file" name="payroll" required accept=".xlsx" id="payrollInput">
                    
                    <label>Billing Report (Excel)</label>
                    <input type="file" name="billing" required accept=".xlsx" id="billingInput">
                    
                    <button type="submit" id="submitBtn">Run Reconciliation</button>
                </form>
            </div>
            
            <div id="progressSection" class="progress-section">
                <h3>Reconciliation Progress</h3>
                <div class="progress-bar">
                    <div id="progressFill" class="progress-fill" style="width: 0%;">0%</div>
                </div>
                
                <div class="progress-info">
                    <div class="info-item">
                        <div class="info-label">ELAPSED TIME</div>
                        <div class="info-value" id="elapsedTime">0s</div>
                    </div>
                    <div class="info-item">
                        <div class="info-label">ESTIMATED REMAINING</div>
                        <div class="info-value" id="estimatedTime">--</div>
                    </div>
                </div>
                
                <div class="status-message">
                    <strong>Stage:</strong> <span id="stageText">Initializing...</span><br>
                    <strong>Status:</strong> <span id="statusText">Preparing files...</span>
                </div>
            </div>
            
            <div id="errorSection" class="error-section"></div>
            <div id="successSection" class="success-section"></div>
        </div>
        
        <script>
            let progressInterval = null;
            
            function formatTime(seconds) {
                if (!seconds || seconds < 0) return '--';
                if (seconds < 60) return seconds + 's';
                return Math.floor(seconds / 60) + 'm ' + (seconds % 60) + 's';
            }
            
            async function updateProgress() {
                try {
                    const response = await fetch('/progress');
                    if (!response.ok) return;
                    
                    const data = await response.json();
                    if (!data) return;
                    
                    if (data.status === 'processing') {
                        const progress = Math.min(data.progress || 0, 100);
                        document.getElementById('progressFill').style.width = progress + '%';
                        document.getElementById('progressFill').textContent = progress + '%';
                        document.getElementById('stageText').textContent = data.stage || 'Processing...';
                        document.getElementById('statusText').textContent = data.message || 'Working...';
                        document.getElementById('elapsedTime').textContent = formatTime(data.elapsed_time);
                        document.getElementById('estimatedTime').textContent = formatTime(data.estimated_time);
                    }
                } catch (error) {
                    console.log('Progress update error:', error);
                }
            }
            
            document.getElementById('reconcileForm').addEventListener('submit', async (e) => {
                e.preventDefault();
                
                document.getElementById('progressSection').classList.add('active');
                document.getElementById('errorSection').classList.remove('active');
                document.getElementById('successSection').classList.remove('active');
                document.getElementById('submitBtn').disabled = true;
                
                const formData = new FormData();
                formData.append('payroll', document.getElementById('payrollInput').files[0]);
                formData.append('billing', document.getElementById('billingInput').files[0]);
                
                progressInterval = setInterval(updateProgress, 500);
                updateProgress();
                
                try {
                    const response = await fetch('/reconcile', {
                        method: 'POST',
                        body: formData
                    });
                    
                    if (response.ok) {
                        const blob = await response.blob();
                        const url = window.URL.createObjectURL(blob);
                        const a = document.createElement('a');
                        a.href = url;
                        a.download = `reconciliation_${new Date().toISOString().split('T')[0]}.xlsx`;
                        a.click();
                        
                        document.getElementById('successSection').textContent = '✓ Success! Your report has been downloaded.';
                        document.getElementById('successSection').classList.add('active');
                        document.getElementById('progressSection').classList.remove('active');
                    } else {
                        const errorText = await response.text();
                        showError('Reconciliation failed: ' + errorText);
                    }
                } catch (error) {
                    showError('Network error: ' + error.message);
                } finally {
                    if (progressInterval) clearInterval(progressInterval);
                    document.getElementById('submitBtn').disabled = false;
                }
            });
            
            function showError(message) {
                document.getElementById('errorSection').textContent = '✗ Error: ' + message;
                document.getElementById('errorSection').classList.add('active');
                document.getElementById('progressSection').classList.remove('active');
            }
        </script>
    </body>
    </html>
    '''


@app.route('/progress', methods=['GET'])
def progress():
    """Get current progress"""
    try:
        return jsonify(current_progress), 200
    except Exception as e:
        logger.error(f"Progress error: {e}")
        return jsonify({'status': 'error', 'progress': 0}), 200


@app.route('/reconcile', methods=['POST'])
def reconcile():
    """Reconciliation endpoint"""
    try:
        global start_time
        start_time = None
        
        # Get files
        payroll_file = request.files.get('payroll')
        billing_file = request.files.get('billing')
        
        if not payroll_file or not billing_file:
            return 'Missing files', 400
        
        logger.info(f"Received: {payroll_file.filename}, {billing_file.filename}")
        update_progress('File Loading', 5, 'Loading payroll...')
        
        # Import pandas here to catch import errors
        try:
            import pandas as pd
            logger.info("Pandas imported successfully")
        except ImportError as e:
            logger.error(f"Pandas import error: {e}")
            return f'Pandas import failed: {e}', 500
        
        # Load files
        try:
            payroll_df = pd.read_excel(payroll_file)
            logger.info(f"Loaded {len(payroll_df)} payroll rows")
            update_progress('File Loading', 15, 'Loading billing...')
            
            billing_df = pd.read_excel(billing_file)
            logger.info(f"Loaded {len(billing_df)} billing rows")
            update_progress('File Loading', 25, 'Files loaded')
            
        except Exception as e:
            logger.error(f"Excel load error: {e}")
            return f'Error reading files: {e}', 400
        
        # Import reconciliation
        try:
            from reconciliation import BillPayReconciler
            logger.info("Reconciliation module imported successfully")
        except ImportError as e:
            logger.error(f"Reconciliation import error: {e}")
            return f'Reconciliation module not found: {e}', 500
        except Exception as e:
            logger.error(f"Reconciliation import error: {e}")
            return f'Error loading reconciliation: {e}', 500
        
        # Run reconciliation
        try:
            logger.info("Starting reconciliation...")
            update_progress('Reconciliation', 30, 'Processing...')
            
            reconciler = BillPayReconciler(payroll_df, billing_df)
            results = reconciler.reconcile()
            
            logger.info(f"Reconciliation complete: {results['summary']['total_matches']} matches")
            update_progress('Reconciliation', 85, 'Creating Excel report...')
            
        except Exception as e:
            logger.error(f"Reconciliation error: {e}", exc_info=True)
            return f'Reconciliation failed: {e}', 500
        
        # Generate Excel
        try:
            from openpyxl import Workbook
            from openpyxl.styles import PatternFill, Font
            
            logger.info("Generating Excel...")
            update_progress('Excel Generation', 90, 'Creating workbook...')
            
            wb = Workbook()
            wb.remove(wb.active)
            
            # Summary tab
            ws = wb.create_sheet('Summary', 0)
            summary = results['summary']
            
            ws['A1'] = 'Bill vs Pay Reconciliation Report'
            ws['A1'].font = Font(size=14, bold=True, color="FFFFFF")
            ws['A1'].fill = PatternFill(start_color="1F5F99", end_color="1F5F99", fill_type="solid")
            ws.merge_cells('A1:B1')
            
            ws['A3'] = 'Total Matches'
            ws['B3'] = summary['total_matches']
            ws['A4'] = 'Match Rate %'
            ws['B4'] = f"{summary['match_rate_payroll']:.2f}%"
            ws['A5'] = 'Matched Amount'
            ws['B5'] = f"${summary['matched_payroll_amount']:,.2f}"
            
            ws.column_dimensions['A'].width = 25
            ws.column_dimensions['B'].width = 25
            
            # Matches tab
            if len(results['matches']) > 0:
                ws_m = wb.create_sheet('Matches')
                matches_df = results['matches']
                
                for col_idx, col_name in enumerate(matches_df.columns, 1):
                    cell = ws_m.cell(row=1, column=col_idx)
                    cell.value = col_name
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill(start_color="1F5F99", end_color="1F5F99", fill_type="solid")
                
                for row_idx, row_data in enumerate(matches_df.itertuples(), 2):
                    for col_idx, value in enumerate(row_data[1:], 1):
                        ws_m.cell(row=row_idx, column=col_idx).value = value
            
            # Save file
            output_path = f'/tmp/reconciliation_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx'
            wb.save(output_path)
            logger.info(f"Excel saved: {output_path}")
            update_progress('Complete', 100, 'Done!')
            
            return send_file(
                output_path,
                as_attachment=True,
                download_name=f"reconciliation_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
                mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
        
        except Exception as e:
            logger.error(f"Excel generation error: {e}", exc_info=True)
            return f'Error generating Excel: {e}', 500
    
    except Exception as e:
        logger.error(f"Reconcile endpoint error: {e}", exc_info=True)
        return f'Unexpected error: {e}', 500


@app.route('/health', methods=['GET'])
def health():
    """Health check"""
    return jsonify({'status': 'healthy'}), 200


@app.errorhandler(404)
def not_found(error):
    return 'Not found', 404


@app.errorhandler(500)
def server_error(error):
    logger.error(f"Server error: {error}")
    return 'Server error', 500


if __name__ == '__main__':
    host = os.environ.get('HOST', '127.0.0.1')
    port = int(os.environ.get('PORT', 8765))
    
    logger.info(f"Starting server on {host}:{port}")
    app.run(host=host, port=port, debug=False, threaded=True)
